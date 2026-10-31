from openpyxl import load_workbook
from decimal import Decimal, InvalidOperation
from pathlib import Path
import csv
import sys

SOURCE = Path("/home/tchadmin/HBP Master  PMJAY MHIS 6.xlsx")
OUTPUT = Path("/tmp/tch_mhis_procedures_import.csv")

if not SOURCE.is_file():
    sys.exit(f"Workbook not found: {SOURCE}")

ws = load_workbook(
    SOURCE,
    read_only=True,
    data_only=True
)["Procedure master - IPD"]

fields = [
    "source_row_number",
    "source_serial_number",
    "procedure_code",
    "specialty",
    "medical_surgical",
    "package_name",
    "procedure_name",
    "procedure_price",
    "source_price_text",
    "average_length_of_stay",
    "stratification_options",
    "stratification_rate",
    "implants",
    "preauth_documents",
    "claim_documents",
]

def clean(value):
    return "" if value is None else str(value).strip()

records = []
seen_rows = set()

for excel_row, row in enumerate(
    ws.iter_rows(min_row=2, values_only=True),
    start=2
):
    if not any(value is not None for value in row):
        continue

    code = clean(row[10])

    if not code:
        sys.exit(f"Missing procedure code at Excel row {excel_row}")

    price_text = clean(row[5])
    price = ""

    if price_text:
        try:
            amount = Decimal(price_text)
            if not amount.is_finite() or amount < 0:
                raise InvalidOperation
            price = f"{amount:.2f}"
        except (InvalidOperation, ValueError):
            pass

    serial = clean(row[0])

    if not serial.isdigit():
        sys.exit(f"Invalid serial number at Excel row {excel_row}")

    if excel_row in seen_rows:
        sys.exit(f"Duplicate source row {excel_row}")

    seen_rows.add(excel_row)

    records.append([
        excel_row,
        int(serial),
        code,
        clean(row[1]),
        clean(row[2]),
        clean(row[3]),
        clean(row[4]),
        price,
        price_text,
        clean(row[6]),
        clean(row[7]),
        clean(row[8]),
        clean(row[9]),
        clean(row[11]),
        clean(row[12]),
    ])

if len(records) != 3232:
    sys.exit(
        f"Expected 3232 procedure rows, found {len(records)}. "
        "Import preparation aborted."
    )

with OUTPUT.open("w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(fields)
    writer.writerows(records)

print("IMPORT FILE PREPARED")
print("Total records:", len(records))
print("Output:", OUTPUT)
print("Database unchanged.")
